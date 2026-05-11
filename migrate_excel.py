import openpyxl
from database import SessionLocal, Snapshot, Ranking, TeamMember
from datetime import datetime, date
import re

def parse_jade(jade_val):
    if jade_val is None or jade_val == "—":
        return None
    if isinstance(jade_val, (int, float)):
        return int(jade_val)
    s = str(jade_val).replace(',', '').upper()
    multiplier = 1
    if 'K' in s:
        multiplier = 1000
        s = s.replace('K', '')
    elif 'M' in s:
        multiplier = 1000000
        s = s.replace('M', '')
    try:
        return int(float(s) * multiplier)
    except:
        return None

def parse_rank(rank_val):
    if rank_val is None or rank_val == "—":
        return None
    if isinstance(rank_val, int):
        return rank_val
    s = str(rank_val).replace('#', '')
    try:
        return int(s)
    except:
        return None

def build_username_map(db):
    """
    Map alias from Excel -> real username in DB.
    Source 1: rankings from snapshot 1 (29/04 - real JSON import)
    Source 2: team_members table
    """
    # 1. rankings ngày 29/04
    latest_snapshot = db.query(Snapshot).filter(Snapshot.date == date(2026, 4, 29)).first()
    alias_to_username = {}
    
    if latest_snapshot:
        latest_rankings = db.query(Ranking).filter(Ranking.snapshot_id == latest_snapshot.id).all()
        for r in latest_rankings:
            if r.alias:
                alias_to_username[r.alias] = r.username
            alias_to_username[r.username] = r.username
            
    # 2. team_members
    team_members = db.query(TeamMember).all()
    for tm in team_members:
        if tm.alias:
            alias_to_username[tm.alias] = tm.username
        alias_to_username[tm.username] = tm.username
        
    return alias_to_username

def migrate():
    wb = openpyxl.load_workbook('ARO_Sprint2_v24_28Apr.xlsx', data_only=True)
    sheet = wb['📊 Master Tracker']
    db = SessionLocal()
    
    # BƯỚC 1: Xóa snapshots cũ migrate sai (note chứa "Migrated")
    # Giữ lại snapshot ngày 29/04 (ID 1 hoặc note khác)
    old_snapshots = db.query(Snapshot).filter(Snapshot.note.like("Migrated%")).all()
    for s in old_snapshots:
        db.delete(s)
    db.commit()
    print(f"Deleted {len(old_snapshots)} old migrated snapshots")
    
    # BƯỚC 2: Parse dates từ row 4
    dates_info = [] # (date, is_delayed)
    for col in range(2, 20):
        raw_val = sheet.cell(row=4, column=col).value
        if raw_val is None:
            dates_info.append((None, False))
            continue
        
        is_delayed = '*' in str(raw_val)
        val_clean = str(raw_val).replace('*', '').strip()
        try:
            d = datetime.strptime(f"{val_clean}-2026", "%d-%m-%Y").date()
            dates_info.append((d, is_delayed))
        except:
            dates_info.append((None, False))
            
    # BƯỚC 3: Collect aliases
    all_aliases = []
    for row in range(5, 150):
        alias = sheet.cell(row=row, column=1).value
        if alias and str(alias).strip():
            all_aliases.append((row, str(alias).strip()))
    print(f"Found {len(all_aliases)} users in Excel")
    
    # BƯỚC 4: Build username map
    alias_to_username = build_username_map(db)
    
    # BƯỚC 5: Migrate
    for i, (d, is_delayed) in enumerate(dates_info):
        if d is None: continue
        
        rank_col = 2 + i
        jade_col = 20 + i
        
        snapshot = db.query(Snapshot).filter(Snapshot.date == d).first()
        if not snapshot:
            snapshot = Snapshot(
                date=d,
                is_delayed=is_delayed,
                note="Migrated from Master Tracker"
            )
            db.add(snapshot)
            db.flush()
            
        rankings_to_add = []
        for row, alias in all_aliases:
            rank_val = sheet.cell(row=row, column=rank_col).value
            jade_val = sheet.cell(row=row, column=jade_col).value
            
            rank = parse_rank(rank_val)
            jade = parse_jade(jade_val)
            
            if rank is None and jade is None:
                continue
                
            username = alias_to_username.get(alias, alias)
            
            # Check duplicate in this snapshot
            existing = db.query(Ranking).filter(Ranking.snapshot_id == snapshot.id, Ranking.username == username).first()
            if existing: continue
            
            ranking = Ranking(
                snapshot_id=snapshot.id,
                rank=rank,
                username=username,
                alias=alias,
                jade=jade or 0,
                t1_refs=0,
                t2_refs=0
            )
            rankings_to_add.append(ranking)
            
        if rankings_to_add:
            db.add_all(rankings_to_add)
            db.flush()
            print(f"  {d}: {len(rankings_to_add)} users")
            
    db.commit()
    db.close()
    print("Migration complete. Run refresh_stats.py next.")

if __name__ == "__main__":
    migrate()
