#!/usr/bin/env python3
"""
ARO Sprint 2 — Report Generator (ARO_report.py)
=================================================
Chạy ARO_update.py rồi đọc DASHBOARD và in báo cáo Markdown hoàn chỉnh.
Claude gọi script này để tạo báo cáo hàng ngày.

Usage:
    python ARO_report.py                    # Dùng ngày trong INPUT sheet
    python ARO_report.py --date 2026-04-29  # Chỉ định ngày
    python ARO_report.py --json data.json   # Dùng JSON file
    python ARO_report.py --read-only        # Chỉ đọc DASHBOARD, không update
"""

import os, sys, argparse, subprocess
from datetime import date, datetime, timedelta
from openpyxl import load_workbook

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR    = r"C:\Users\Administrator\OneDrive\Documents\ARO S2 top 100 tracking"
MASTER_FILE = os.path.join(BASE_DIR, "ARO_MASTER.xlsx")
UPDATE_SCRIPT = os.path.join(BASE_DIR, "ARO_update.py")

LEADER   = "nau***@gmail.com"
EVENT_END = date(2026, 5, 31)

def val(cell):
    """Get cell value as string, strip whitespace."""
    v = cell.value
    if v is None:
        return "—"
    return str(v).strip()

def run_update(args_extra=None):
    """Run ARO_update.py and print its output."""
    cmd = [sys.executable, UPDATE_SCRIPT]
    if args_extra:
        cmd.extend(args_extra)
    print("[INFO] Running ARO_update.py...")
    result = subprocess.run(cmd, capture_output=True, text=True, cwd=BASE_DIR, encoding="utf-8")
    if result.stdout:
        print(result.stdout)
    if result.returncode != 0:
        print(f"[ERROR] Update failed:\n{result.stderr}")
        sys.exit(1)

def read_dashboard(wb):
    """Extract structured data from DASHBOARD sheet."""
    ws = wb["🏆 DASHBOARD"]
    data = {}

    # ── Section 1: nau*** stats (rows 4-14) ──
    stat_labels = {4: "rank", 5: "jade", 6: "t1", 7: "t2",
                   8: "delta", 9: "wrate", 10: "proj",
                   11: "prize", 12: "gap2", 13: "gap3", 14: "date"}
    data["leader"] = {}
    for row, key in stat_labels.items():
        data["leader"][key] = val(ws.cell(row, 4))  # merged B-J, value in col B (index 2) or D

    # Retry: value might be in col 2 (B) after merge
    for row, key in stat_labels.items():
        # Try column 2 (B) and 4 (D)
        for col in [2, 4]:
            v = ws.cell(row, col).value
            if v and str(v).strip() not in ("", "—", label):
                data["leader"][key] = str(v).strip()
                break
        # label is col 1
        label = val(ws.cell(row, 1))

    # Re-read more carefully
    leader_data = {}
    for row, key in stat_labels.items():
        # Section 1 layout: A=label, B:J=merged value
        for col in [2, 4, 3]:
            v = ws.cell(row, col).value
            if v is not None and str(v).strip():
                leader_data[key] = str(v).strip()
                break
        if key not in leader_data:
            leader_data[key] = "—"
    data["leader"] = leader_data

    # ── Section 2: Top 10 (rows 18-27) ──
    data["top10"] = []
    for row in range(18, 28):
        rank_val = ws.cell(row, 1).value
        if rank_val is None:
            continue
        row_data = [val(ws.cell(row, col)) for col in range(1, 12)]
        if any(v not in ("", "—", "None") for v in row_data[1:]):
            data["top10"].append(row_data)

    # ── Section 3: Top T1 refs (rows 32-41) ──
    data["top_t1"] = []
    for row in range(32, 42):
        rank_val = ws.cell(row, 1).value
        if rank_val is None:
            continue
        row_data = [val(ws.cell(row, col)) for col in range(1, 9)]
        if any(v not in ("", "—", "None") for v in row_data[1:]):
            data["top_t1"].append(row_data)

    # ── Section 4: Team (rows 46-60) ──
    data["team"] = []
    for row in range(46, 61):
        uname = ws.cell(row, 1).value
        if not uname or str(uname).strip() in ("", "None"):
            break
        row_data = [val(ws.cell(row, col)) for col in range(1, 10)]
        data["team"].append(row_data)

    # ── Section 5: Thresholds (rows 64-65) ──
    data["thresholds"] = []
    for row in [64, 65]:
        label = val(ws.cell(row, 1))
        jade  = val(ws.cell(row, 2))
        if label != "—":
            data["thresholds"].append((label, jade))

    # ── Section 6: Alerts (rows 72-79) ──
    data["alerts"] = []
    for row in range(72, 80):
        status = ws.cell(row, 1).value
        detail = ws.cell(row, 2).value
        if status and detail:
            data["alerts"].append((str(status).strip(), str(detail).strip()))

    # ── Title for date ──
    title_cell = ws["A1"].value or ""
    data["title"] = str(title_cell)

    return data


def format_report(data: dict, today: date) -> str:
    """Format structured data into a Markdown report."""
    ldr   = data.get("leader", {})
    top10 = data.get("top10", [])
    top_t1= data.get("top_t1", [])
    team  = data.get("team", [])
    thresh= data.get("thresholds", [])
    alerts= data.get("alerts", [])

    date_str   = today.strftime("%d/%m/%Y")
    days_left  = max(0, (EVENT_END - today).days)

    # Build headline
    jade_today  = ldr.get("jade", "?")
    delta_today = ldr.get("delta", "?")
    rank_today  = ldr.get("rank", "#?")

    # Determine notable headline
    headline = f"nau*** giữ vững {rank_today} — {jade_today} jade"
    if "gap2" in ldr and ldr["gap2"] not in ("—", ""):
        headline += f" | Lead +{ldr['gap2']}"

    lines = []
    lines.append(f"## 📊 Báo Cáo {date_str} — {headline}")
    lines.append(f"> 🕐 Còn {days_left} ngày đến 31/05/2026\n")

    # ── Cảnh báo nổi bật (nếu có) ─────────────────────────
    critical = [a for a in alerts if a[0] in ("🔴", "🟡")]
    if critical:
        for status, detail in critical:
            lines.append(f"> {status} **CẢNH BÁO:** {detail}\n")

    # ── Báo Cáo 1: Top 10 ─────────────────────────────────
    lines.append("### 🥇 Báo Cáo 1 — Top 10 BXH\n")
    lines.append("| # | Username | Jade | Δ Hôm nay | Δ Hôm qua | W.Rate | Proj 31/05 | Prize | T1 | T2 | Note |")
    lines.append("|---|----------|------|-----------|-----------|--------|-----------|-------|----|----|------|")
    for row in top10:
        if len(row) >= 11:
            lines.append(f"| {' | '.join(row[:11])} |")

    lines.append("")

    # ── Báo Cáo 2: Top T1 Refs ────────────────────────────
    lines.append("### 🔗 Báo Cáo 2 — Top 10 T1 Refs Nhiều Nhất\n")
    lines.append("| # | Username | T1 | T2 | Jade | Δ/ngày | W.Rate | Verdict |")
    lines.append("|---|----------|----|----|------|--------|--------|---------|")
    for row in top_t1:
        if len(row) >= 8:
            lines.append(f"| {' | '.join(row[:8])} |")

    lines.append("")

    # ── Báo Cáo 3: Team nau*** ────────────────────────────
    lines.append("### 👥 Báo Cáo 3 — Team nau***\n")

    # Leader stats box
    lines.append(f"**nau*** (LEADER)**")
    lines.append(f"- Rank: **{ldr.get('rank','?')}** | Jade: **{ldr.get('jade','?')}** | T1: {ldr.get('t1','?')} | T2: {ldr.get('t2','?')}")
    lines.append(f"- Δ hôm nay: **{ldr.get('delta','?')}** | W.Rate: {ldr.get('wrate','?')}/ngày")
    lines.append(f"- Dự báo 31/05: **{ldr.get('proj','?')}** | Prize: **{ldr.get('prize','?')}**")
    lines.append(f"- Gap vs #2: {ldr.get('gap2','?')} | Gap vs #3: {ldr.get('gap3','?')}")
    lines.append("")

    lines.append("| Username | Level | Rank | Jade | Δ Hôm nay | W.Rate | Proj 31/05 | Prize | Ghi chú |")
    lines.append("|----------|-------|------|------|-----------|--------|-----------|-------|---------|")
    for row in team:
        if len(row) >= 9:
            lines.append(f"| {' | '.join(row[:9])} |")

    lines.append("")

    # ── Ngưỡng & Cảnh báo ────────────────────────────────
    lines.append("### 📏 Ngưỡng & Cảnh báo\n")
    for label, jade in thresh:
        lines.append(f"- **Jade {label}:** {jade}")

    lines.append("")
    if alerts:
        lines.append("**Tất cả cảnh báo:**")
        for status, detail in alerts:
            lines.append(f"- {status} {detail}")
    else:
        lines.append("✅ Không có cảnh báo đặc biệt hôm nay.")

    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description="ARO Sprint 2 Report")
    parser.add_argument("--date",      help="Override date (YYYY-MM-DD)")
    parser.add_argument("--json",      help="Path to JSON file with BXH data")
    parser.add_argument("--auto",      action="store_true",
                        help="Tự tìm file bxh-*.json mới nhất")
    parser.add_argument("--read-only", action="store_true",
                        help="Only read DASHBOARD, skip update")
    parser.add_argument("--delayed",   action="store_true",
                        help="Mark today as delayed")
    args = parser.parse_args()

    # Determine today
    if args.date:
        today = datetime.strptime(args.date, "%Y-%m-%d").date()
    else:
        today = date.today()

    # Run update (unless read-only)
    if not args.read_only:
        update_args = []
        if args.date:
            update_args += ["--date", args.date]
        if args.json:
            update_args += ["--json", args.json]
        if args.auto:
            update_args += ["--auto"]
        if args.delayed:
            update_args += ["--delayed"]
        run_update(update_args)

    # Read DASHBOARD
    if not os.path.exists(MASTER_FILE):
        print(f"[ERROR] {MASTER_FILE} not found!")
        sys.exit(1)

    wb   = load_workbook(MASTER_FILE, data_only=True)
    data = read_dashboard(wb)

    # Generate report
    report = format_report(data, today)

    print("\n" + "=" * 60)
    print("REPORT OUTPUT:")
    print("=" * 60)
    print(report)


if __name__ == "__main__":
    main()
